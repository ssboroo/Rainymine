"""Bounded XLSX/STL exchange and read-only canonical SQL Server views.

Geobank installations have site-specific schemas. A DBA must expose the documented
Rainymine* views. This adapter never infers the vendor schema or writes to its DB.
"""
import csv
import re
import zipfile
from copy import deepcopy
from pathlib import Path
from .core import finite
from .data_tools import SCHEMAS, rows_to_values
from .geometry import validate_mesh_data


def xlsx_to_csv(source,destination,sheet):
    from openpyxl import load_workbook
    source=Path(source)
    if source.stat().st_size>32_000_000:raise ValueError('XLSX хэмжээ 32 MB-аас их.')
    with zipfile.ZipFile(source) as archive:
        if sum(x.file_size for x in archive.infolist())>128_000_000:
            raise ValueError('XLSX задалсан хэмжээ 128 MB-аас их.')
    workbook=load_workbook(source,read_only=True,data_only=False,keep_links=False)
    try:
        if sheet not in workbook.sheetnames:raise ValueError('Excel sheet олдсонгүй.')
        ws=workbook[sheet]
        if ws.max_column>256 or ws.max_row>100001:raise ValueError('Excel 256 багана / 100000 өгөгдлийн мөрөөс их.')
        with open(destination,'w',encoding='utf-8-sig',newline='') as target:
            writer=csv.writer(target)
            for index,row in enumerate(ws.iter_rows()):
                if index>100000:raise ValueError('Excel мөрийн тоо хэт их.')
                values=[]
                for cell in row:
                    if cell.data_type=='f':raise ValueError('Excel formula байна. Values болгон экспортлоно уу; formula-г тооцохгүй.')
                    if cell.data_type=='e':raise ValueError('Excel error cell байна.')
                    values.append('' if cell.value is None else str(cell.value))
                if any(v!='' for v in values):writer.writerow(values)
    finally:workbook.close()


def read_stl(path):
    import struct
    data=Path(path).read_bytes()
    if len(data)>32_000_000:raise ValueError('STL хэмжээ 32 MB-аас их.')
    vertices=[];faces=[];ids={}
    def triangle(points):
        face=[]
        for point in points:
            finite(*point)
            if point not in ids:ids[point]=len(vertices);vertices.append(list(point))
            face.append(ids[point])
        faces.append(face)
    count=struct.unpack_from('<I',data,80)[0] if len(data)>=84 else 0
    if len(data)==84+50*count:
        if count>400000:raise ValueError('STL face тоо хэт их.')
        for i in range(count):
            record=struct.unpack_from('<12fH',data,84+i*50)
            triangle([tuple(record[j:j+3]) for j in (3,6,9)])
    else:
        try:lines=[line.strip() for line in data.decode('ascii').splitlines() if line.strip()]
        except UnicodeDecodeError:raise ValueError('STL binary урт / ASCII encoding буруу.') from None
        if not lines or not lines[0].startswith('solid') or not lines[-1].startswith('endsolid'):
            raise ValueError('STL solid boundary алга.')
        if (len(lines)-2)%7:raise ValueError('STL facet бүтэц буруу.')
        for i in range(1,len(lines)-1,7):
            part=lines[i:i+7]
            if not part[0].startswith('facet normal ') or part[1]!='outer loop' or part[5]!='endloop' or part[6]!='endfacet':
                raise ValueError('STL facet бүтэц буруу.')
            points=[]
            for line in part[2:5]:
                words=line.split()
                if len(words)!=4 or words[0]!='vertex':raise ValueError('STL vertex буруу.')
                points.append(tuple(float(x) for x in words[1:]))
            triangle(points)
    mesh={'name':Path(path).stem,'role':'imported_stl','vertices':vertices,'faces':faces}
    validate_mesh_data(mesh);return mesh


def write_stl(path,mesh):
    # ASCII with Python double coordinates avoids vtk float-point truncation at UTM scales.
    import numpy as np
    from .core import atomic_write
    validate_mesh_data(mesh);v=np.asarray(mesh['vertices'],dtype=float);lines=['solid RAINYMINE']
    for a,b,c in mesh['faces']:
        normal=np.cross(v[b]-v[a],v[c]-v[a]);norm=float(np.linalg.norm(normal))
        if norm==0:raise ValueError('STL-д талбайгүй гурвалжин байна.')
        normal/=norm;lines.append('facet normal '+' '.join(format(x,'.17g') for x in normal));lines.append('outer loop')
        for i in (a,b,c):lines.append('vertex '+' '.join(format(x,'.17g') for x in v[i]))
        lines.extend(['endloop','endfacet'])
    lines.append('endsolid RAINYMINE');atomic_write(path,'\n'.join(lines)+'\n')


VIEW_NAMES={'collars':'RainymineCollars','surveys':'RainymineSurveys',
            'assays':'RainymineAssays','lithology':'RainymineLithology'}


def database_snapshot(project,dsn,schema='dbo',connector=None):
    """Read four canonical views atomically using integrated-auth DSN.

    Credentials are managed by Windows/ODBC, never stored in the project. Requires
    a SELECT-only database account and a driver supporting snapshot transactions.
    """
    if not re.fullmatch(r'[A-Za-z0-9 _.-]{1,100}',dsn) or not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]{0,63}',schema):
        raise ValueError('DSN / SQL schema нэр буруу.')
    if connector is None:
        try:import pyodbc
        except ImportError:raise ValueError('ODBC adapter суулгаагүй. Windows build ашиглана уу.') from None
        connector=pyodbc.connect
    connection=None
    try:
        connection=connector('DSN='+dsn+';Trusted_Connection=yes;',timeout=10,autocommit=False)
        connection.timeout=20
        cursor=connection.cursor()
        # Fail rather than quietly mixing tables from different database states.
        cursor.execute('SET TRANSACTION ISOLATION LEVEL SNAPSHOT')
        cursor.execute('BEGIN TRANSACTION')
        candidate=deepcopy(project)
        for kind,view in VIEW_NAMES.items():
            fields=SCHEMAS[kind]
            sql='SELECT TOP (100001) '+','.join('['+f+']' for f in fields)+' FROM ['+schema+'].['+view+']'
            cursor.execute(sql);raw=cursor.fetchmany(100001)
            if len(raw)>100000:raise ValueError(view+' 100000 мөрөөс их.')
            parsed=[]
            for row in raw:
                if len(row)!=len(fields):raise ValueError(view+' баганын тоо буруу.')
                item={}
                for field,value in zip(fields,row):
                    if value is None:raise ValueError(view+': NULL '+field)
                    if field in ('hole_id','rock'):
                        if not isinstance(value,str) or not value.strip():raise ValueError(view+': text '+field)
                        item[field]=value.strip()
                    else:
                        item[field]=float(value);finite(item[field])
                parsed.append(item)
            setattr(candidate,kind,rows_to_values(kind,parsed))
        candidate.workspace.setdefault('imports',[]).append({'source':'canonical SQL Server views',
                    'counts':{k:len(getattr(candidate,k)) for k in VIEW_NAMES},'snapshot_isolation':True})
        candidate.validate()
        # Old derived geology/estimation data must not look current after replacing inputs.
        candidate.workspace.pop('block_model',None);candidate.workspace.pop('block_slice_positions',None)
        candidate.workspace.pop('pit_optimization',None)
        for mesh in candidate.workspace.get('meshes',[]):
            if mesh.get('role') in ('implicit_grade','optimized_pit'):mesh['name']='[ӨМНӨХ ӨГӨГДӨЛ] '+mesh['name']
        candidate.validate();return candidate
    except ValueError:raise
    except Exception:
        # Drivers often echo connection strings in errors; show a sanitized diagnostic.
        raise ValueError('ODBC уншилт амжилтгүй. DSN, SELECT эрх, canonical views, SNAPSHOT тохиргоог DBA-тай шалгана уу.') from None
    finally:
        if connection is not None:
            try:connection.rollback()
            except Exception:pass
            try:connection.close()
            except Exception:pass
